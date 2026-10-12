import unittest
import io
from docx import Document

from utils.resume_parser import clean_extracted_text, extract_txt, extract_docx
from utils.information_extractor import (
    extract_email,
    extract_phone,
    extract_linkedin,
    extract_github,
    extract_portfolio,
    extract_location,
    extract_name_heuristic,
    estimate_experience_years,
    extract_all_contact_info
)
from utils.skill_extractor import (
    extract_skills_offline,
    extract_categorized_skills,
    extract_skills,
    canonicalize_skill
)
from utils.ats_score import (
    calculate_ats_score_tfidf,
    calculate_offline_ats_audit,
    clean_job_description_boilerplate
)
from utils.export_helper import (
    create_cover_letter_docx,
    create_ats_report_markdown,
    create_ats_report_docx,
    create_interview_prep_docx,
    create_outreach_docx,
    create_roadmap_docx,
    create_roadmap_markdown,
    create_linkedin_docx,
    SAMPLE_PROFILES
)
from utils.offline_heuristics import (
    generate_offline_bullet_rewrite,
    generate_offline_tailored_summary,
    generate_offline_cover_letter,
    generate_offline_cold_outreach,
    generate_offline_linkedin_profile,
    generate_offline_interview_questions,
    evaluate_offline_interview_answer,
    generate_offline_followup_question,
    generate_offline_interview_cheat_sheet,
    generate_offline_skill_gap_roadmap
)
from utils.gemini_service import (
    clean_json_response,
    SUPPORTED_MODELS,
    DEFAULT_MODEL,
    save_api_key_to_env,
    remove_saved_api_key,
    get_active_api_key,
    rewrite_resume_bullet,
    generate_tailored_summary,
    generate_cover_letter,
    generate_cold_outreach,
    generate_linkedin_profile_content,
    generate_interview_questions,
    evaluate_interview_answer,
    generate_skill_gap_roadmap,
    parse_full_resume_with_gemini
)


class TestCareerAISuite(unittest.TestCase):

    def setUp(self):
        self.sample = SAMPLE_PROFILES["AI / Machine Learning Engineer"]
        self.sample_fs = SAMPLE_PROFILES["Full-Stack Software Engineer"]

    def test_text_cleaning(self):
        raw = "Hello\x00 World\xa0Test\r\n\r\n\r\n\r\nEnd"
        cleaned = clean_extracted_text(raw)
        self.assertNotIn("\x00", cleaned)
        self.assertNotIn("\xa0", cleaned)
        self.assertIn("Hello World Test\n\nEnd", cleaned)

    def test_contact_extractor(self):
        text = self.sample["resume_text"]
        contact = extract_all_contact_info(text)
        self.assertEqual(contact["full_name"], "Alex Chen")
        self.assertEqual(contact["email"], "alex.chen.ai@example.com")
        self.assertIn("415", contact["phone"])
        self.assertEqual(contact["location"], "San Francisco, CA")
        self.assertIn("linkedin.com/in/alexchen-ai", contact["linkedin"])
        self.assertIn("github.com/alexchen-ml", contact["github"])
        self.assertEqual(contact["portfolio"], "https://alexchen.dev")
        self.assertGreaterEqual(contact["estimated_exp_years"], 4)

    def test_docx_table_extraction(self):
        # Create an in-memory DOCX with tables and paragraphs
        doc = Document()
        doc.add_paragraph("Header Title: Alex Chen")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Skill Category"
        table.rows[0].cells[1].text = "Skills"
        table.rows[1].cells[0].text = "Languages"
        table.rows[1].cells[1].text = "Python, SQL, C++"
        
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        
        extracted = extract_docx(bio)
        self.assertIn("Header Title: Alex Chen", extracted)
        self.assertIn("Python, SQL, C++", extracted)

    def test_skill_canonicalization(self):
        self.assertEqual(canonicalize_skill("golang"), "Go")
        self.assertEqual(canonicalize_skill("k8s"), "Kubernetes")
        self.assertEqual(canonicalize_skill("reactjs"), "React")
        self.assertEqual(canonicalize_skill("postgres"), "PostgreSQL")
        self.assertEqual(canonicalize_skill("sklearn"), "Scikit-learn")

    def test_offline_skill_extraction(self):
        text = "I build apps in Python, FastAPI, Docker, and PostgreSQL with React and AWS."
        skills = extract_skills_offline(text)
        self.assertIn("Python", skills)
        self.assertIn("FastAPI", skills)
        self.assertIn("Docker", skills)
        self.assertIn("PostgreSQL", skills)
        self.assertIn("React", skills)
        self.assertIn("AWS", skills)

    def test_skill_categorization(self):
        text = self.sample["resume_text"]
        categorized = extract_categorized_skills(text)
        self.assertIn("Programming Languages", categorized)
        self.assertIn("Frameworks & Libraries", categorized)
        self.assertIn("Cloud, DevOps & Tools", categorized)
        self.assertIn("Databases & Storage", categorized)
        self.assertIn("Python", categorized["Programming Languages"])
        self.assertIn("Docker", categorized["Cloud, DevOps & Tools"])

    def test_ats_scoring_offline(self):
        resume = self.sample["resume_text"]
        jd = self.sample["job_desc"]
        
        audit = calculate_offline_ats_audit(resume, jd)
        self.assertGreater(audit["score_overall"], 0)
        self.assertGreater(audit["score_skills"], 0)
        self.assertIn("keyword_analysis", audit)
        self.assertIn("matched_keywords", audit["keyword_analysis"])
        self.assertTrue(len(audit["keyword_analysis"]["matched_keywords"]) > 0)
        self.assertIn("Python", audit["keyword_analysis"]["matched_keywords"])
        self.assertTrue(len(audit["section_audit"]) >= 5)

    def test_json_cleaner(self):
        markdown_json = "```json\n{\"score\": 95, \"rating\": \"Excellent\"}\n```"
        parsed = clean_json_response(markdown_json)
        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["score"], 95)

        trailing_comma_json = "{\"items\": [1, 2, ], }"
        parsed_fixed = clean_json_response(trailing_comma_json)
        self.assertIsNotNone(parsed_fixed)
        self.assertEqual(parsed_fixed["items"], [1, 2])

    def test_export_helpers(self):
        # 1. Cover letter docx
        cl_buf = create_cover_letter_docx("Dear Hiring Team,\n\nI am thrilled to apply.")
        self.assertGreater(len(cl_buf.getvalue()), 1000)

        # 2. ATS Report Markdown
        audit = calculate_offline_ats_audit(self.sample["resume_text"], self.sample["job_desc"])
        md = create_ats_report_markdown(audit, candidate_name="Alex Chen")
        self.assertIn("Alex Chen", md)
        self.assertIn("Overall ATS Compatibility", md)

        # 3. ATS Report Docx
        ats_docx = create_ats_report_docx(audit, candidate_name="Alex Chen")
        self.assertGreater(len(ats_docx.getvalue()), 1000)

        # 4. Interview Prep Docx
        mock_qs = [{"id": 1, "type": "Technical", "question": "Explain RAG architecture", "interviewer_intent": "Depth"}]
        mock_fb = {1: {"rating": "Outstanding", "score_total": 92, "star_breakdown": {"situation_score": 9}, "strengths": ["Clear metrics"], "growth_areas": [], "sample_answer": "Model ans"}}
        prep_docx = create_interview_prep_docx(mock_qs, mock_fb, candidate_name="Alex Chen")
        self.assertGreater(len(prep_docx.getvalue()), 1000)

        # 5. Outreach Docx
        outreach_data = [{"type": "LinkedIn Connection Note", "subject": "Connect", "message": "Hi, saw your post", "tip": "Keep it under 300 chars"}]
        outreach_docx = create_outreach_docx(outreach_data, candidate_name="Alex Chen")
        self.assertGreater(len(outreach_docx.getvalue()), 1000)

    def test_txt_and_empty_edge_cases(self):
        # TXT parser
        txt_content = "John Doe\njohn@example.com\nPython, Docker"
        extracted = extract_txt(txt_content)
        self.assertEqual(extracted, txt_content)

        # Empty inputs
        self.assertEqual(extract_email(""), "Not Found")
        self.assertEqual(extract_phone(""), "Not Found")
        self.assertEqual(calculate_ats_score_tfidf("", ""), 0.0)

    def test_boilerplate_cleaning(self):
        raw_jd = "Python Engineer needed.\n\nEqual Opportunity Employer. All applicants considered.\n\nBenefits include 401k."
        cleaned = clean_job_description_boilerplate(raw_jd)
        self.assertNotIn("Equal Opportunity Employer", cleaned)
        self.assertNotIn("Benefits include 401k", cleaned)
        self.assertIn("Python Engineer needed.", cleaned)

    def test_api_key_persistence(self):
        test_env_filename = ".env.test_tmp"
        test_key = "AIzaSyTEST_KEY_123456789"

        try:
            # 1. Save key
            saved = save_api_key_to_env(test_key, env_filename=test_env_filename)
            self.assertTrue(saved)

            # 2. Check active key
            active = get_active_api_key()
            self.assertEqual(active, test_key)

            # 3. Clear key
            cleared = remove_saved_api_key(env_filename=test_env_filename)
            self.assertTrue(cleared)
        finally:
            import os
            if os.path.exists(test_env_filename):
                os.remove(test_env_filename)

    def test_navigation_logic(self):
        from app import PAGES, set_page, render_step_header, render_step_navigation
        import streamlit as st

        # Initial state
        set_page(PAGES[0])
        self.assertEqual(st.session_state.current_page, PAGES[0])
        self.assertEqual(st.session_state.nav_page, PAGES[0])

        # Test set_page to page 2 (index 1) and page 7 (index 6)
        set_page(PAGES[1])
        self.assertEqual(st.session_state.current_page, PAGES[1])
        self.assertEqual(st.session_state.nav_page, PAGES[1])

        set_page(PAGES[6])
        self.assertEqual(st.session_state.current_page, PAGES[6])
        self.assertEqual(st.session_state.nav_page, PAGES[6])

        # Reset back to step 1
        set_page(PAGES[0])
        self.assertEqual(st.session_state.current_page, PAGES[0])

    def test_offline_heuristics_suite(self):
        resume = self.sample["resume_text"]
        jd = self.sample["job_desc"]

        # 1. Bullet rewrite
        raw_bullet = "Responsible for managing the team's database and writing queries."
        rew_res = generate_offline_bullet_rewrite(raw_bullet, job_context="PostgreSQL, Query Optimization")
        self.assertTrue(rew_res["is_offline"])
        self.assertIn("variations", rew_res)
        self.assertEqual(len(rew_res["variations"]), 3)
        self.assertTrue(any("38%" in v["rewritten"] for v in rew_res["variations"]))

        # 2. Tailored summary
        summary = generate_offline_tailored_summary(resume, jd)
        self.assertIn("experience", summary.lower())
        self.assertGreater(len(summary), 80)

        # 3. Cover letter
        cl = generate_offline_cover_letter(resume, jd, company_name="Acme Corp", hiring_manager="Jane Doe")
        self.assertIn("Jane Doe", cl)
        self.assertIn("Acme Corp", cl)
        self.assertIn("Alex Chen", cl)

        # 4. Cold outreach
        li_note = generate_offline_cold_outreach(resume, jd, outreach_type="LinkedIn Connection Note (<300 chars)", company_name="Acme Corp")
        self.assertTrue(li_note["is_offline"])
        self.assertLessEqual(len(li_note["message_body"]), 300)

        email_pitch = generate_offline_cold_outreach(resume, jd, outreach_type="Direct Hiring Manager Cold Email Pitch", company_name="Acme Corp")
        self.assertIn("Alex Chen", email_pitch["subject_or_headline"])
        self.assertIn("Alex Chen", email_pitch["message_body"])

        # 5. LinkedIn profile
        li_prof = generate_offline_linkedin_profile(resume, target_role="AI Engineer")
        self.assertEqual(len(li_prof["headlines"]), 3)
        self.assertIn("AI Engineer", li_prof["about_summary"])
        self.assertEqual(len(li_prof["top_skills_to_pin"]), 5)

        # 6. Interview questions
        qs = generate_offline_interview_questions(resume, jd, seniority_level="Senior Engineer")
        self.assertEqual(len(qs), 5)
        self.assertEqual(qs[0]["type"], "Technical")

        # 7. Answer evaluation
        eval_res = evaluate_offline_interview_answer(
            qs[0]["question"],
            "In my previous role, I reduced latency by 45% by implementing Redis caching and indexing PostgreSQL.",
            jd,
            resume
        )
        self.assertGreaterEqual(eval_res["score_total"], 40)
        self.assertIn("star_breakdown", eval_res)
        self.assertTrue(len(eval_res["strengths"]) > 0)

        # 8. Follow-up
        followup = generate_offline_followup_question(qs[0]["question"], "I used redis caching", jd)
        self.assertIn("cache", followup.lower())

        # 9. Cheat sheet
        cs = generate_offline_interview_cheat_sheet(resume, jd)
        self.assertEqual(len(cs["anchor_stories"]), 3)
        self.assertEqual(len(cs["curveball_questions_to_watch"]), 3)
        self.assertEqual(len(cs["questions_to_ask_interviewer"]), 4)

        # 10. Roadmap
        rm = generate_offline_skill_gap_roadmap(["Docker", "Kubernetes", "AWS"], target_role="Cloud Engineer")
        self.assertEqual(len(rm["weekly_plan"]), 4)
        self.assertEqual(len(rm["recommended_projects"]), 2)

    def test_service_fallbacks_without_api_key(self):
        resume = self.sample["resume_text"]
        jd = self.sample["job_desc"]

        # All service functions should fall back seamlessly to offline heuristics without throwing or erroring
        profile = parse_full_resume_with_gemini(resume, api_key="")
        self.assertIsNotNone(profile)
        self.assertEqual(profile["full_name"], "Alex Chen")

        rew = rewrite_resume_bullet("Worked on api endpoints", job_context="FastAPI", api_key="")
        self.assertIn("variations", rew)

        summ = generate_tailored_summary(resume, jd, api_key="")
        self.assertGreater(len(summ), 50)

        letter = generate_cover_letter(resume, jd, company_name="Tech Co", api_key="")
        self.assertIn("Tech Co", letter)

        outreach = generate_cold_outreach(resume, jd, api_key="")
        self.assertIn("message_body", outreach)

        qs = generate_interview_questions(resume, jd, api_key="")
        self.assertEqual(len(qs), 5)

        eval_res = evaluate_interview_answer("Tell me about a project", "I built a model that reduced costs by 20%", jd, resume, api_key="")
        self.assertIn("score_total", eval_res)

        rm = generate_skill_gap_roadmap(["GraphQL", "Terraform"], api_key="")
        self.assertIn("weekly_plan", rm)

    def test_new_document_exports(self):
        # 1. Roadmap DOCX & Markdown
        rm = generate_offline_skill_gap_roadmap(["FastAPI", "Docker"], target_role="Backend Engineer")
        docx_buf = create_roadmap_docx(rm, candidate_name="Alex Chen", target_role="Backend Engineer")
        self.assertGreater(len(docx_buf.getvalue()), 1000)

        rm_md = create_roadmap_markdown(rm, candidate_name="Alex Chen", target_role="Backend Engineer")
        self.assertIn("30-Day Upskilling", rm_md)
        self.assertIn("Week 1", rm_md)

        # 2. LinkedIn DOCX
        li_prof = generate_offline_linkedin_profile(self.sample["resume_text"], target_role="AI Engineer")
        li_docx = create_linkedin_docx(li_prof, candidate_name="Alex Chen", target_role="AI Engineer")
        self.assertGreater(len(li_docx.getvalue()), 1000)

    def test_explicit_experience_extraction(self):
        text_with_explicit = "Senior Machine Learning Engineer with 6+ years of experience building LLM pipelines."
        years = estimate_experience_years(text_with_explicit)
        self.assertEqual(years, 6)


if __name__ == "__main__":
    unittest.main()

